# -*- coding: utf-8 -*-
import io
import os
from datetime import date

import openpyxl

import campos

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RUTA_PLANTILLA = os.path.join(BASE_DIR, 'plantilla.xlsx')

MESES = ['ENE', 'FEB', 'MAR', 'ABR', 'MAY', 'JUN',
         'JUL', 'AGO', 'SEP', 'OCT', 'NOV', 'DIC']
DIAS = ['LUNES', 'MARTES', 'MIERCOLES', 'JUEVES', 'VIERNES', 'SABADO', 'DOMINGO']


def texto_fecha(iso):
    d = date.fromisoformat(iso)
    return f'{d.day} DE {MESES[d.month - 1]} DE {d.year}'


def nombre_archivo(iso):
    d = date.fromisoformat(iso)
    return f'REPORTE {DIAS[d.weekday()]} {d.day} DE {MESES[d.month - 1]} {d.year}.xlsx'


def generar(reg):
    if not os.path.exists(RUTA_PLANTILLA):
        raise RuntimeError(
            'No se encontró plantilla.xlsx junto a la aplicación; '
            'no es posible generar el Excel.')
    wb = openpyxl.load_workbook(RUTA_PLANTILLA)
    ws = wb['LENIN']
    ws['B3'] = texto_fecha(reg['fecha'])
    if reg.get('clima'):
        ws['C3'] = reg['clima']
    if reg.get('predicador'):
        ws['D3'] = reg['predicador']
    if reg.get('tema'):
        ws['E3'] = reg['tema']
    for col, celda in campos.EXCEL_MAP.items():
        valor = reg.get(col)
        if valor is not None:
            ws[celda] = valor
    try:
        wb.calculation.fullCalcOnLoad = True
    except AttributeError:
        pass
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, nombre_archivo(reg['fecha'])
